import openpyxl
import os
import json
import re
from pathlib import Path

def convert_curriculum_semester_2_2569():
    base_dir = Path(__file__).resolve().parent.parent
    config_path = base_dir / "data" / "timetable_config.json"
    backup_path = base_dir / "data" / "timetable_config.backup_before_2_2569.json"

    plan_dir = Path("/Users/allarmmac/Downloads/แผนการเรียน")
    path_pvc = plan_dir / "รหัส 69 ปวช.  หลักสูตร67 .xlsx"
    path_pvs = plan_dir / "รหัส 69 ปวส. (ทวิ) หลักสูตร67 .xlsx"
    path_m6  = plan_dir / "รหัส 69 ปวส. ม.6 (ทวิ) หลักสูตร67 .xlsx"

    # Backup
    if config_path.exists():
        with open(config_path, "r", encoding="utf-8") as f:
            old_config = json.load(f)
        with open(backup_path, "w", encoding="utf-8") as f:
            json.dump(old_config, f, ensure_ascii=False, indent=2)
    else:
        old_config = {}

    teachers = old_config.get("teachers", [])
    if not any(t["id"] == "T_UNASSIGNED" for t in teachers):
        teachers.insert(0, {
            "id": "T_UNASSIGNED",
            "name": "(ยังไม่ระบุครูผู้สอน)",
            "max_periods_per_day": 12,
            "max_periods_per_week": 35,
            "is_head": False,
            "qualification": "รอมอบหมาย",
            "special_duty": "รอมอบหมาย"
        })

    rooms = old_config.get("rooms", [])

    # Defined groups for semester 2/2569
    groups = [
        {
            "id": "G_CHO_1_1",
            "name": "ชอ.1/1 (ปวช.1 อิเล็กทรอนิกส์)",
            "level": "VOC_CERT",
            "student_count": 25,
            "pvs_18_weeks": True,
            "is_internship": False
        },
        {
            "id": "G_PVS_1",
            "name": "สอ.1/1 (ปวส.1 ทวิภาคี)",
            "level": "HIGH_VOC_CERT",
            "student_count": 20,
            "pvs_18_weeks": True,
            "is_internship": False
        },
        {
            "id": "G_PVS_M6_1",
            "name": "สอ.1/2 (ปวส.1 ม.6 ทวิ)",
            "level": "HIGH_VOC_CERT",
            "student_count": 20,
            "pvs_18_weeks": True,
            "is_internship": False
        }
    ]

    def determine_room_type(name, p):
        if p == 0:
            return "CLASSROOM"
        if any(k in name for k in ["ภาษา", "คณิต", "สังคม", "วิทยา"]):
            return "CLASSROOM"
        if any(k in name for k in ["เชื่อม", "เครื่องมือกล", "ฝึกฝีมือ", "นิวเมติกส์"]):
            return "LAB_ENGINE"
        return "LAB_ELECTRIC"

    sources = [
        (path_pvc, "G_CHO_1_1", "PVC"),
        (path_pvs, "G_PVS_1", "PVS"),
        (path_m6, "G_PVS_M6_1", "PVSM6")
    ]

    courses_dict = {}
    assignments = []
    ass_counter = 1

    for filepath, grp_id, prefix in sources:
        if not filepath.exists():
            print(f"Warning: file not found: {filepath}")
            continue
        wb = openpyxl.load_workbook(filepath, data_only=True)
        if "2.2569" not in wb.sheetnames:
            print(f"Warning: sheet 2.2569 not found in {filepath.name}")
            continue

        ws = wb["2.2569"]
        for r in ws.iter_rows(values_only=True):
            if not r or len(r) < 5: continue
            code = str(r[0]).strip() if r[0] is not None else ""
            name = str(r[1]).strip() if r[1] is not None else ""
            t, p, c = r[2], r[3], r[4]
            if any(char.isdigit() for char in code) and "-" in code:
                try:
                    t_int = int(t) if t is not None else 0
                    p_int = int(p) if p is not None else 0
                    c_int = int(c) if c is not None else 0
                    tot = t_int + p_int
                except:
                    continue

                clean_code = re.sub(r'[^0-9A-Za-z]', '_', code)
                cid = f"C_{clean_code}"
                rtype = determine_room_type(name, p_int)
                is_rot = (grp_id in ["G_PVS_1", "G_PVS_M6_1"] and tot >= 5)
                if is_rot:
                    ctype = "ROTATION_BASE"
                elif p_int == 0:
                    ctype = "THEORY"
                else:
                    ctype = "PRACTICE"

                # Special case for 6 periods: split into 2 sessions of 3 periods so it fits around lunch
                if tot == 6:
                    c_item = {
                        "id": cid,
                        "name": name,
                        "code": code,
                        "course_type": ctype,
                        "periods_per_session": 3,
                        "sessions_per_week": 2,
                        "required_room_type": rtype,
                        "allow_merge": False,
                        "base_id": None
                    }
                    courses_dict[cid] = c_item

                    assignments.append({
                        "id": f"L_{ass_counter}_{clean_code}_p1",
                        "course_id": cid,
                        "primary_group_id": grp_id,
                        "teacher_id": "T_UNASSIGNED",
                        "secondary_teacher_id": None,
                        "secondary_group_id": None,
                        "is_rotation": False,
                        "teaching_mode": "SINGLE"
                    })
                    ass_counter += 1
                    assignments.append({
                        "id": f"L_{ass_counter}_{clean_code}_p2",
                        "course_id": cid,
                        "primary_group_id": grp_id,
                        "teacher_id": "T_UNASSIGNED",
                        "secondary_teacher_id": None,
                        "secondary_group_id": None,
                        "is_rotation": False,
                        "teaching_mode": "SINGLE"
                    })
                    ass_counter += 1
                else:
                    c_item = {
                        "id": cid,
                        "name": name,
                        "code": code,
                        "course_type": ctype,
                        "periods_per_session": tot,
                        "sessions_per_week": 1,
                        "required_room_type": rtype,
                        "allow_merge": False,
                        "base_id": f"BASE_{clean_code}" if is_rot else None
                    }
                    courses_dict[cid] = c_item

                    assignments.append({
                        "id": f"L_{ass_counter}_{clean_code}",
                        "course_id": cid,
                        "primary_group_id": grp_id,
                        "teacher_id": "T_UNASSIGNED",
                        "secondary_teacher_id": None,
                        "secondary_group_id": None,
                        "is_rotation": is_rot,
                        "teaching_mode": "SINGLE"
                    })
                    ass_counter += 1

    new_config = {
        "teachers": teachers,
        "rooms": rooms,
        "groups": groups,
        "courses": courses_dict,
        "assignments": assignments
    }

    with open(config_path, "w", encoding="utf-8") as f:
        json.dump(new_config, f, ensure_ascii=False, indent=2)

    print("Successfully converted and saved ภาคเรียนที่ 2/2569!")
    print(f"Total groups: {len(groups)}")
    print(f"Total unique courses: {len(courses_dict)}")
    print(f"Total lesson assignments: {len(assignments)}")

if __name__ == "__main__":
    convert_curriculum_semester_2_2569()

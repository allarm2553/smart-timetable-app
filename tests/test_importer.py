import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import io
import base64
import openpyxl
from fastapi.testclient import TestClient
from src.api.main import app
from src.api.data_manager import data_manager
from src.api.importer import BulkDataImporter

client = TestClient(app)

def test_is_valid_course_code():
    # Valid course codes
    assert BulkDataImporter.is_valid_course_code("20000-1102") is True
    assert BulkDataImporter.is_valid_course_code("30000-1201") is True
    assert BulkDataImporter.is_valid_course_code("20101-2001") is True
    assert BulkDataImporter.is_valid_course_code("2101-2001") is True
    assert BulkDataImporter.is_valid_course_code("ว30101") is True
    assert BulkDataImporter.is_valid_course_code("ENG 101") is True
    assert BulkDataImporter.is_valid_course_code("20000-2001") is True

    # Invalid items (category headers, section numbers, totals, summaries, headers)
    assert BulkDataImporter.is_valid_course_code("1. หมวดวิชาสมรรถนะแกนกลาง ( 4 หน่วยกิต)") is False
    assert BulkDataImporter.is_valid_course_code("1.1 กลุ่มสมรรถนะภาษาและการสื่อสาร") is False
    assert BulkDataImporter.is_valid_course_code("2. หมวดวิชาสมรรถนะวิชาชีพ") is False
    assert BulkDataImporter.is_valid_course_code("รวมสัปดาห์ละ") is False
    assert BulkDataImporter.is_valid_course_code("รวม") is False
    assert BulkDataImporter.is_valid_course_code("รวมทั้งสิ้น") is False
    assert BulkDataImporter.is_valid_course_code("1.1") is False
    assert BulkDataImporter.is_valid_course_code("2.0") is False
    assert BulkDataImporter.is_valid_course_code("รหัสวิชา") is False
    assert BulkDataImporter.is_valid_course_code("code") is False
    assert BulkDataImporter.is_valid_course_code("") is False
    assert BulkDataImporter.is_valid_course_code(None) is False
    assert BulkDataImporter.is_valid_course_code("   ") is False
    print("✅ test_is_valid_course_code passed")

def test_template_csv():
    csv_bytes = BulkDataImporter.get_template_csv()
    assert len(csv_bytes) > 0
    # Must have UTF-8 BOM
    assert csv_bytes.startswith(b'\xef\xbb\xbf')
    text = csv_bytes.decode('utf-8-sig')
    assert "รหัสวิชา" in text
    assert "งานเครื่องยนต์แก๊สโซลีน" in text
    print("✅ test_template_csv passed")

def test_api_get_template():
    res = client.get("/api/import/template")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment" in res.headers.get("content-disposition", "")
    assert len(res.content) > 100
    print("✅ test_api_get_template passed")

def test_process_import_append():
    data_manager.reset_to_default()
    csv_sample = """รหัสวิชา,ชื่อวิชา,ประเภทวิชา,จำนวนคาบ,ประเภทห้อง,ครูผู้สอน,กลุ่มผู้เรียน,กลุ่มเรียนรวม,หมุนเวียนฐาน
20101-9901,วิชาทดสอบนำเข้าใหม่,practice,4,lab,ครูทดสอบพิเศษ,ชอ.1/1,,ใช่
"""
    b64_content = base64.b64encode(csv_sample.encode("utf-8")).decode("utf-8")
    
    payload = {
        "filename": "test_import.csv",
        "content_base64": b64_content,
        "mode": "append"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["is_success"] is True
    assert data["imported_count"] == 1
    assert data["new_teachers_count"] >= 1  # 'ครูทดสอบพิเศษ' is newly created

    # Verify teacher was added to data manager
    all_d = data_manager.get_all_data()
    assert any(t["name"] == "ครูทดสอบพิเศษ" for t in all_d["teachers"])
    print("✅ test_process_import_append passed")

def test_process_import_filters_non_course_rows():
    data_manager.reset_to_default()
    # CSV containing category headers and total summaries
    csv_with_categories = """รหัสวิชา,ชื่อวิชา,ประเภทวิชา,จำนวนคาบ,ประเภทห้อง,ครูผู้สอน,กลุ่มผู้เรียน,กลุ่มเรียนรวม,หมุนเวียนฐาน
1. หมวดวิชาสมรรถนะแกนกลาง ( 4 หน่วยกิต),,,,,,,,
1.1 กลุ่มสมรรถนะภาษาและการสื่อสาร,,,,,,,,
20000-1102,ภาษาไทยเพื่ออาชีพ,ทฤษฎี,1,ห้องบรรยาย,ครูสมใจ,ชอ.1/1,,ไม่ใช่
20000-1202,ภาษาอังกฤษเพื่อการสื่อสารในงานอาชีพ,ทฤษฎี,2,ห้องบรรยาย,ครูนภา,ชอ.1/1,,ไม่ใช่
2. หมวดวิชาสมรรถนะวิชาชีพ,,,,,,,,
20101-2001,งานเครื่องยนต์แก๊สโซลีน,ปฏิบัติ,5,ห้องปฏิบัติการ,ครูพงษ์สถิต,ชอ.1/1,,ไม่ใช่
รวมสัปดาห์ละ,,,,,,,,
"""
    b64_content = base64.b64encode(csv_with_categories.encode("utf-8")).decode("utf-8")
    payload = {
        "filename": "plan_with_categories.csv",
        "content_base64": b64_content,
        "mode": "replace"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["is_success"] is True
    # Only the 3 actual courses must be imported!
    assert data["imported_count"] == 3
    assert data["skipped_count"] == 4  # 3 category rows + 1 total row

    # Verify that NONE of the category headers or totals became courses
    all_d = data_manager.get_all_data()
    course_codes = [c["code"] for c in all_d["courses"]]
    assert "20000-1102" in course_codes
    assert "20000-1202" in course_codes
    assert "20101-2001" in course_codes
    for c in all_d["courses"]:
        assert "หมวด" not in c["name"]
        assert "กลุ่ม" not in c["name"]
        assert "รวม" not in c["name"]
        assert not c["id"].startswith("C_00") # No dummy generated codes
    print("✅ test_process_import_filters_non_course_rows passed")

def test_process_import_vocational_xlsx():
    data_manager.reset_to_default()
    # Create an in-memory vocational study plan Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2.2569"

    ws.append(["แผนการเรียนหลักสูตรประกาศนียบัตรวิชาชีพ (ปวช.) พุทธศักราช 2567"])
    ws.append(["สาขาวิชาช่างอิเล็กทรอนิกส์"])
    ws.append(["ภาคเรียนที่ 2 ปีการศึกษา 2569"])
    # Header row at row 4
    ws.append(["รหัสวิชา", "รายวิชา", "ท", "ป", "น", "หมายเหตุ"])
    ws.append(["1. หมวดวิชาสมรรถนะแกนกลาง (4 หน่วยกิต)", "", "", "", "", ""])
    ws.append(["1.1 กลุ่มสมรรถนะภาษาและการสื่อสาร", "", "", "", "", ""])
    ws.append(["20000-1102", "ภาษาไทยเพื่ออาชีพ", 1, 0, 1, ""])
    ws.append(["20000-1202", "ภาษาอังกฤษเพื่อการสื่อสารในงานอาชีพ", 1, 1, 1, ""])
    ws.append(["2. หมวดวิชาสมรรถนะวิชาชีพ", "", "", "", "", ""])
    ws.append(["20105-2001", "วงจรไฟฟ้ากระแสสลับ", 1, 4, 3, ""])
    ws.append(["20105-2002", "เครื่องวัดไฟฟ้าและอิเล็กทรอนิกส์", 1, 4, 3, ""])
    ws.append(["รวมสัปดาห์ละ", "", 4, 9, 8, ""])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    b64_content = base64.b64encode(xlsx_bytes).decode("utf-8")
    payload = {
        "filename": "curriculum_plan.xlsx",
        "content_base64": b64_content,
        "mode": "replace"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["is_success"] is True
    # Strictly 4 courses imported
    assert data["imported_count"] == 4
    assert data["skipped_count"] >= 3

    # Check course details
    all_d = data_manager.get_all_data()
    assert len(all_d["courses"]) == 4
    for c in all_d["courses"]:
        # Verify course name is NEVER equal to course code
        assert c["name"] != c["code"]
        assert "20000" not in c["name"] and "20105" not in c["name"]
        if c["code"] == "20000-1102":
            assert c["name"] == "ภาษาไทยเพื่ออาชีพ"
            assert c["periods_per_session"] == 1
            assert c["course_type"] == "THEORY"
        elif c["code"] == "20105-2001":
            assert c["name"] == "วงจรไฟฟ้ากระแสสลับ"
            assert c["periods_per_session"] == 5 # 1 + 4
            assert c["course_type"] in ["PRACTICE", "THEORY_PRACTICE"]

    # Solve should succeed
    solve_res = client.post("/api/solve/current")
    assert solve_res.status_code == 200
    assert solve_res.json()["is_success"] is True
    print("✅ test_process_import_vocational_xlsx passed")

def test_process_import_curriculum_with_title_and_subheaders():
    data_manager.reset_to_default()
    wb = openpyxl.Workbook()
    # Sheet 1: Title with 'รหัส 69' and subheader
    ws1 = wb.active
    ws1.title = "1.2569"
    ws1.append(["แผนการเรียนมุ่งสมรรถนะอาชีพ หมวดวิชา รหัส 69"]) # Title containing 'รหัส'
    ws1.append(["สาขาวิชาช่างยนต์"])
    ws1.append(["ภาคเรียนที่ 1 ปีการศึกษา 2569"])
    ws1.append(["รหัสวิชา", "ภาคเรียนที่ 1/2569", "ท", "ป", "น"]) # Col 2 has semester title
    ws1.append(["", "รายวิชา", "", "", ""]) # Subheader specifies 'รายวิชา'
    ws1.append(["20000-1101", "ภาษาไทยเพื่อสื่อสาร", 1, 0, 1])
    ws1.append(["20000-1201", "ภาษาอังกฤษเพื่อการสื่อสาร", 1, 1, 1])

    # Sheet 2: Side-by-side tables (ใบปะหน้า)
    ws2 = wb.create_sheet("ใบปะหน้า ปวช")
    ws2.append(["ใบปะหน้าแผนการเรียน ปวช. 2567"])
    ws2.append(["สาขาวิชาช่างยนต์"])
    ws2.append([])
    ws2.append(["รหัสวิชา", "รายวิชา", "ท", "ป", "น", "", "รหัสวิชา", "รายวิชา", "ท", "ป", "น"])
    ws2.append(["20101-1001", "งานเครื่องยนต์แก๊สโซลีน", 1, 4, 3, "", "20101-2002", "งานเครื่องยนต์ดีเซล", 1, 4, 3])
    ws2.append(["20101-1002", "งานระบบส่งกำลัง", 1, 4, 3, "", "20101-2003", "งานระบบเบรก", 1, 4, 3])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    b64_content = base64.b64encode(xlsx_bytes).decode("utf-8")
    payload = {
        "filename": "curriculum_multi_sheet.xlsx",
        "content_base64": b64_content,
        "mode": "replace"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["is_success"] is True
    assert data["imported_count"] == 6

    all_d = data_manager.get_all_data()
    course_map = {c["code"]: c["name"] for c in all_d["courses"]}
    
    # Verify accurate Thai names for all courses
    assert course_map["20000-1101"] == "ภาษาไทยเพื่อสื่อสาร"
    assert course_map["20000-1201"] == "ภาษาอังกฤษเพื่อการสื่อสาร"
    assert course_map["20101-1001"] == "งานเครื่องยนต์แก๊สโซลีน"
    assert course_map["20101-2002"] == "งานเครื่องยนต์ดีเซล"
    assert course_map["20101-1002"] == "งานระบบส่งกำลัง"
    assert course_map["20101-2003"] == "งานระบบเบรก"

    for code, name in course_map.items():
        assert name != code, f"Course {code} name should not be code!"
        assert not name.startswith("วิชา 20"), f"Course {code} should have real title, got {name}"

    print("✅ test_process_import_curriculum_with_title_and_subheaders passed")

def test_process_import_replace_and_solve():
    # Use template CSV with replace mode
    template_bytes = BulkDataImporter.get_template_csv()
    b64_content = base64.b64encode(template_bytes).decode("utf-8")
    
    payload = {
        "filename": "curriculum.csv",
        "content_base64": b64_content,
        "mode": "replace"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["is_success"] is True
    assert data["imported_count"] == 5
    assert data["total_assignments"] == 5

    # Run solver on current data
    solve_res = client.post("/api/solve/current")
    assert solve_res.status_code == 200
    s_data = solve_res.json()
    assert s_data["is_success"] is True
    assert s_data["total_lessons_scheduled"] == 5
    print("✅ test_process_import_replace_and_solve passed")

    # Reset back to default
    data_manager.reset_to_default()

def test_theory_practice_course_type():
    data_manager.reset_to_benchmark()
    # 1. Add course with THEORY_PRACTICE
    payload = {
        "code": "20105-2101",
        "name": "ไมโครคอนโทรลเลอร์ประยุกต์",
        "course_type": "THEORY_PRACTICE",
        "periods_per_session": 4,
        "required_room_type": "LAB_ELECTRIC",
        "primary_group_id": "G_CHO_1_1",
        "teacher_id": "T_PONG",
        "teaching_mode": "SINGLE"
    }
    res = client.post("/api/assignments", json=payload)
    assert res.status_code == 200
    res_data = res.json()
    assert res_data["is_success"] is True

    # 2. Verify stored course_type is THEORY_PRACTICE
    all_d = data_manager.get_all_data()
    course = next((c for c in all_d["courses"] if c["code"] == "20105-2101"), None)
    assert course is not None
    assert course["course_type"] == "THEORY_PRACTICE"
    assert course["periods_per_session"] == 4

    # 3. Update assignment and verify editing preserves THEORY_PRACTICE
    ass = next((a for a in all_d["assignments"] if a["course_id"] == course["id"]), None)
    assert ass is not None
    update_payload = {
        "name": "ไมโครคอนโทรลเลอร์ประยุกต์ (แก้ไข)",
        "code": "20105-2101",
        "course_type": "THEORY_PRACTICE",
        "periods_per_session": 3,
        "required_room_type": "LAB_ELECTRIC",
        "primary_group_id": "G_CHO_1_1",
        "teacher_id": "T_PONG"
    }
    up_res = client.put(f"/api/assignments/{ass['id']}", json=update_payload)
    assert up_res.status_code == 200

    all_d = data_manager.get_all_data()
    updated_course = next((c for c in all_d["courses"] if c["id"] == course["id"]), None)
    assert updated_course["course_type"] == "THEORY_PRACTICE"
    assert updated_course["name"] == "ไมโครคอนโทรลเลอร์ประยุกต์ (แก้ไข)"
    assert updated_course["periods_per_session"] == 3

    # 4. Solve and ensure it generates timetable successfully
    solve_res = client.post("/api/solve/current")
    assert solve_res.status_code == 200
    assert solve_res.json()["is_success"] is True

    # Reset back to default
    data_manager.reset_to_default()
    print("✅ test_theory_practice_course_type passed")

def test_inline_group_and_teacher_assignment():
    """Verify that assignments can update primary_group_id and teacher_id via API and on bulk import"""
    data_manager.reset_to_default()
    all_d = data_manager.get_all_data()
    ass = all_d["assignments"][0]
    ass_id = ass["id"]

    # 1. Update primary_group_id
    up_grp = client.put(f"/api/assignments/{ass_id}", json={"primary_group_id": "G_CHO_1_2"})
    assert up_grp.status_code == 200
    assert up_grp.json()["is_success"] is True
    
    # 2. Update teacher_id to another teacher
    teachers = all_d["teachers"]
    other_t = [t["id"] for t in teachers if t["id"] != ass["teacher_id"]][0]
    up_tch = client.put(f"/api/assignments/{ass_id}", json={"teacher_id": other_t})
    assert up_tch.status_code == 200
    
    # Verify in data_manager
    check_d = data_manager.get_all_data()
    updated_ass = next(a for a in check_d["assignments"] if a["id"] == ass_id)
    assert updated_ass["primary_group_id"] == "G_CHO_1_2"
    assert updated_ass["teacher_id"] == other_t

    # 3. Update teacher_id to T_UNASSIGNED
    up_unassigned = client.put(f"/api/assignments/{ass_id}", json={"teacher_id": "T_UNASSIGNED"})
    assert up_unassigned.status_code == 200
    check_d2 = data_manager.get_all_data()
    updated_ass2 = next(a for a in check_d2["assignments"] if a["id"] == ass_id)
    assert updated_ass2["teacher_id"] == "T_UNASSIGNED"

    # 4. Test Bulk Import with target_group_id and default_teacher_id
    csv_content = (
        "รหัสวิชา,ชื่อวิชา,ท,ป,น\n"
        "20101-2001,งานขับเคลื่อนยานยนต์,1,3,2\n"
        "20101-2002,งานเครื่องยนต์แก๊สโซลีน,1,3,2\n"
    )
    res_import = client.post("/api/import/process", json={
        "filename": "curriculum.csv",
        "csv_text": csv_content,
        "mode": "replace",
        "target_group_id": "G_CHO_1_2",
        "default_teacher_id": other_t
    })
    assert res_import.status_code == 200
    imp_data = res_import.json()
    assert imp_data["is_success"] is True
    assert imp_data["imported_count"] == 2

    cur_d = data_manager.get_all_data()
    for a in cur_d["assignments"]:
        assert a["primary_group_id"] == "G_CHO_1_2"
        assert a["teacher_id"] == other_t

    data_manager.reset_to_default()
    print("✅ test_inline_group_and_teacher_assignment passed")

def test_multiple_groups_per_course_with_different_teachers():
    """Verify that a single course can be assigned to multiple groups with different teachers"""
    data_manager.reset_to_default()
    all_d = data_manager.get_all_data()
    course = all_d["courses"][0]

    # Assign same course to Group 1 with Teacher 1
    t1 = all_d["teachers"][0]["id"]
    t2 = all_d["teachers"][1]["id"]
    g1 = all_d["groups"][0]["id"]
    g2 = all_d["groups"][1]["id"]

    res1 = client.post("/api/assignments", json={
        "course_id": course["id"],
        "name": course["name"],
        "code": course["code"],
        "course_type": course["course_type"],
        "periods_per_session": course["periods_per_session"],
        "required_room_type": course["required_room_type"],
        "primary_group_id": g1,
        "teacher_id": t1
    })
    assert res1.status_code == 200
    a1 = res1.json()["data"]["assignment"]
    assert a1["primary_group_id"] == g1
    assert a1["teacher_id"] == t1
    assert a1["course_id"] == course["id"]

    # Assign same course to Group 2 with Teacher 2
    res2 = client.post("/api/assignments", json={
        "course_id": course["id"],
        "name": course["name"],
        "code": course["code"],
        "course_type": course["course_type"],
        "periods_per_session": course["periods_per_session"],
        "required_room_type": course["required_room_type"],
        "primary_group_id": g2,
        "teacher_id": t2
    })
    assert res2.status_code == 200
    a2 = res2.json()["data"]["assignment"]
    assert a2["primary_group_id"] == g2
    assert a2["teacher_id"] == t2
    assert a2["course_id"] == course["id"]
    assert a1["id"] != a2["id"]

    # Verify Timetable Solver successfully schedules both
    solve_res = client.post("/api/solve/current")
    assert solve_res.status_code == 200
    assert solve_res.json()["is_success"] is True

    data_manager.reset_to_default()
    print("✅ test_multiple_groups_per_course_with_different_teachers passed")

def test_process_import_sorts_by_course_code():
    data_manager.reset_to_default()
    csv_scrambled = """รหัสวิชา,ชื่อวิชา,ประเภทวิชา,จำนวนคาบ,ประเภทห้อง,ครูผู้สอน,กลุ่มผู้เรียน,กลุ่มเรียนรวม,หมุนเวียนฐาน
30101-2001,วิชา C ปวส,practice,4,lab,อ. สมคิด,ชอ.1/1,,
20000-1101,วิชา A พื้นฐาน,theory,2,classroom,อ. สมชาย,ชอ.1/1,,
20101-2001,วิชา B ช่างยนต์,practice,3,lab,อ. สมศักดิ์,ชอ.1/1,,
"""
    b64_content = base64.b64encode(csv_scrambled.encode("utf-8")).decode("utf-8")
    payload = {
        "filename": "scrambled_courses.csv",
        "content_base64": b64_content,
        "mode": "replace"
    }
    res = client.post("/api/import/process", json=payload)
    assert res.status_code == 200

    all_data = data_manager.get_all_data()
    course_map = {c["id"]: c for c in all_data["courses"]}
    imported_codes = [course_map[a["course_id"]]["code"] for a in all_data["assignments"]]

    assert imported_codes == ["20000-1101", "20101-2001", "30101-2001"], f"Assignments should be sorted by code, got: {imported_codes}"

    # Test sorting API endpoint
    res_sort = client.post("/api/assignments/sort")
    assert res_sort.status_code == 200
    assert res_sort.json()["is_success"] is True
    print("✅ test_process_import_sorts_by_course_code passed")

if __name__ == "__main__":
    orig_data = data_manager.get_all_data()
    try:
        test_is_valid_course_code()
        test_template_csv()
        test_api_get_template()
        test_process_import_append()
        test_process_import_filters_non_course_rows()
        test_process_import_vocational_xlsx()
        test_process_import_curriculum_with_title_and_subheaders()
        test_process_import_replace_and_solve()
        test_theory_practice_course_type()
        test_inline_group_and_teacher_assignment()
        test_multiple_groups_per_course_with_different_teachers()
        test_process_import_sorts_by_course_code()
        print("\n🎉 ALL BULK IMPORTER TESTS PASSED SUCCESSFULLY! 🎉")
    finally:
        data_manager.teachers = orig_data["teachers"]
        data_manager.rooms = orig_data["rooms"]
        data_manager.groups = orig_data["groups"]
        data_manager.courses = {c["id"]: c for c in orig_data["courses"]}
        data_manager.assignments = orig_data["assignments"]
        data_manager._save()

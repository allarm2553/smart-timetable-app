"""
Vocational Timetable Bulk Data Importer
Supports CSV (UTF-8, UTF-8-SIG, TIS-620, CP874) and native XLSX (via openpyxl)
Strictly imports only valid course items with genuine course codes (skipping category headers, titles, and summaries).
"""
import io
import csv
import re
import openpyxl
from typing import Dict, Any, List, Tuple, Optional

from src.solver.models import CourseType, RoomType, EducationLevel

class BulkDataImporter:
    TEMPLATE_HEADERS = [
        "รหัสวิชา",
        "ชื่อวิชา",
        "ประเภทวิชา",
        "จำนวนคาบ",
        "ประเภทห้อง",
        "ครูผู้สอน",
        "กลุ่มผู้เรียน",
        "กลุ่มเรียนรวม",
        "หมุนเวียนฐาน"
    ]

    SAMPLE_ROWS = [
        ["20101-2001", "งานเครื่องยนต์แก๊สโซลีน", "ปฏิบัติ", "4", "ห้องปฏิบัติการ", "พงษ์สถิต", "ชอ.1/1", "", "ใช่"],
        ["20101-2002", "งานเครื่องยนต์ดีเซล", "ปฏิบัติ", "4", "ห้องปฏิบัติการ", "อ.ป้อ", "ชอ.1/2", "", "ใช่"],
        ["20101-2003", "งานส่งกำลังรถยนต์", "ปฏิบัติ", "4", "ห้องปฏิบัติการ", "อ.สมชาย", "ชอ.1/3", "", "ใช่"],
        ["20000-1201", "ภาษาอังกฤษเพื่อการสื่อสาร", "ทฤษฎี", "2", "ห้องบรรยาย", "ครูนภา", "ชอ.1/1", "ชอ.1/2", "ไม่ใช่"],
        ["20101-2004", "ทฤษฎีเครื่องยนต์แก๊สโซลีน", "ทฤษฎี", "2", "ห้องทั่วไป", "พงษ์สถิต", "ชอ.1/1", "", "ไม่ใช่"]
    ]

    @classmethod
    def is_valid_course_code(cls, code: Any) -> bool:
        """
        Validates whether a string represents an actual vocational course code
        (e.g., 20000-1102, 30000-1201, 20101-2001, 2101-2001, ว30101, ENG101).
        Rejects category headers, section titles, totals, and empty cells.
        """
        if code is None:
            return False
        c = str(code).strip()
        if not c or len(c) < 3 or len(c) > 30:
            return False

        # Must contain at least one digit
        if not any(ch.isdigit() for ch in c):
            return False

        lower_c = c.lower()
        invalid_keywords = [
            "รวม", "หมวด", "กลุ่ม", "ภาคเรียน", "รหัส", "code", "total", "sum",
            "หมายเหตุ", "หน่วยกิต", "ชั่วโมง", "สัปดาห์", "ทฤษฎี", "ปฏิบัติ",
            "รายวิชา", "ชื่อวิชา", "แผนการ", "ระดับ", "สาขา", "ห้อง", "ชั้นปี",
            "กิจกรรม", "คำอธิบาย"
        ]
        if any(kw in lower_c for kw in invalid_keywords):
            return False

        # Reject purely numbering formats like "1.", "1.1", "2.1.3", "3.0"
        if re.match(r'^\d+(\.\d+)*\.?$', c):
            return False

        # Valid course codes do not have multi-word descriptive titles
        if len(c.split()) > 2:
            return False

        return True

    @classmethod
    def get_template_csv(cls) -> bytes:
        """Returns CSV bytes with UTF-8 BOM so Excel opens Thai properly."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(cls.TEMPLATE_HEADERS)
        for r in cls.SAMPLE_ROWS:
            writer.writerow(r)
        
        # Prepend UTF-8 BOM (\ufeff)
        csv_text = "\ufeff" + output.getvalue()
        return csv_text.encode("utf-8")

    @classmethod
    def parse_file_to_rows(cls, file_bytes: bytes, filename: str) -> List[Dict[str, str]]:
        """Parses CSV or XLSX into a list of row dictionaries."""
        filename_lower = filename.lower()
        if filename_lower.endswith(".xlsx") or filename_lower.endswith(".xls"):
            return cls._parse_xlsx(file_bytes)
        else:
            return cls._parse_csv(file_bytes)

    @classmethod
    def _parse_csv(cls, file_bytes: bytes) -> List[Dict[str, str]]:
        # Detect encoding
        text = None
        for enc in ["utf-8-sig", "utf-8", "cp874", "tis-620", "latin1"]:
            try:
                text = file_bytes.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if not text:
            raise ValueError("ไม่สามารถอ่านการเข้ารหัสของไฟล์ CSV ได้ (กรุณาบันทึกเป็น UTF-8)")

        # Sniff delimiter
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            return []

        delimiter = ','
        first_line = lines[0]
        if '\t' in first_line:
            delimiter = '\t'
        elif ';' in first_line and ',' not in first_line:
            delimiter = ';'

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        raw_rows = list(reader)
        if not raw_rows:
            return []

        headers = [h.strip().replace('\ufeff', '') for h in raw_rows[0]]
        dict_rows = []
        for r in raw_rows[1:]:
            if not any(cell.strip() for cell in r):
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                val = r[idx].strip() if idx < len(r) else ""
                row_dict[h] = val
            dict_rows.append(row_dict)

        return dict_rows

    @classmethod
    def _parse_xlsx(cls, file_bytes: bytes) -> List[Dict[str, str]]:
        """
        Reads XLSX using openpyxl, automatically detecting header rows across sheets
        and extracting study plan columns.
        """
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            dict_rows = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if not ws.max_row or ws.max_row < 1:
                    continue

                header_row_idx = None
                col_map = {}

                max_r = min(25, ws.max_row)
                max_c = min(30, ws.max_column if ws.max_column else 10)

                # Scan up to row 25 to find the table header row
                for r_idx in range(1, max_r + 1):
                    row_vals = [ws.cell(r_idx, c).value for c in range(1, max_c + 1)]
                    str_vals = [str(v).strip() if v is not None else "" for v in row_vals]

                    has_code = any("รหัส" in v.lower() or "code" in v.lower() for v in str_vals)
                    has_name = any("ชื่อ" in v.lower() or "รายวิชา" in v.lower() or "name" in v.lower() for v in str_vals)

                    if has_code or (has_name and any(v in ["ท", "ป", "น", "หน่วยกิต", "คาบ", "periods"] for v in str_vals)):
                        header_row_idx = r_idx
                        for c_idx, val in enumerate(str_vals, start=1):
                            if val:
                                col_map[c_idx] = val
                        break

                if header_row_idx is None:
                    # Fallback to row 1
                    header_row_idx = 1
                    for c_idx in range(1, max_c + 1):
                        val = ws.cell(1, c_idx).value
                        if val is not None and str(val).strip():
                            col_map[c_idx] = str(val).strip()

                if not col_map:
                    continue

                # Read following data rows
                for r_idx in range(header_row_idx + 1, ws.max_row + 1):
                    row_dict = {}
                    has_any_val = False
                    for c_idx, h_name in col_map.items():
                        val = ws.cell(r_idx, c_idx).value
                        v_str = str(val).strip() if val is not None else ""
                        if v_str.endswith(".0") and v_str[:-2].isdigit():
                            v_str = v_str[:-2]
                        if v_str:
                            has_any_val = True
                        row_dict[h_name] = v_str

                    if has_any_val:
                        row_dict["_sheet_name"] = sheet_name
                        dict_rows.append(row_dict)

            return dict_rows
        except Exception as e:
            raise ValueError(f"เกิดข้อผิดพลาดในการเปิดไฟล์ Excel: {str(e)}")

    @classmethod
    def _find_field(cls, row: Dict[str, str], aliases: List[str]) -> str:
        """
        Looks up a field using a list of aliases.
        Performs exact match first, then substring match for aliases of len >= 2
        to prevent single-letter aliases ('ท', 'ป', 'น') from matching words like 'ประเภทวิชา'.
        """
        # Pass 1: exact match
        for k, v in row.items():
            if k.startswith("_"):
                continue
            k_clean = k.strip().lower()
            for a in aliases:
                if k_clean == a.lower():
                    return v.strip()

        # Pass 2: substring match (alias must be in header key, len >= 2)
        for k, v in row.items():
            if k.startswith("_"):
                continue
            k_clean = k.strip().lower()
            for a in aliases:
                if len(a) >= 2 and a.lower() in k_clean:
                    return v.strip()

        return ""

    @classmethod
    def _determine_room_type(cls, name: str, course_type: str, raw_room: str) -> str:
        """Infers appropriate vocational room type."""
        r = raw_room.lower() if raw_room else ""
        if any(w in r for w in ["เชื่อม", "เครื่องกล", "กลโรงงาน", "ยานยนต์", "lab_engine"]):
            return RoomType.LAB_ENGINE.value
        if any(w in r for w in ["ไฟฟ้า", "อิเล็ก", "โทรคม", "iot", "lab_electric"]):
            return RoomType.LAB_ELECTRIC.value
        if any(w in r for w in ["ปฏิบัติ", "lab", "ศูนย์", "โรงฝึก"]):
            return RoomType.LAB_ENGINE.value
        if any(w in r for w in ["บรรยาย", "lecture"]):
            return RoomType.LECTURE_HALL.value
        if any(w in r for w in ["ทั่วไป", "classroom", "ห้องเรียน"]):
            return RoomType.CLASSROOM.value

        # Infer from name
        if course_type == CourseType.THEORY.value:
            return RoomType.CLASSROOM.value
        if any(k in name for k in ["ภาษา", "คณิต", "สังคม", "วิทยา", "ศีลธรรม", "พลเมือง", "สุขภาวะ", "ผู้ประกอบการ", "ธุรกิจ"]):
            return RoomType.CLASSROOM.value
        if any(k in name for k in ["เชื่อม", "เครื่องมือกล", "ฝึกฝีมือ", "นิวเมติกส์", "ยานยนต์", "เครื่องกล", "ปรับอากาศ"]):
            return RoomType.LAB_ENGINE.value
        if any(k in name for k in ["อิเล็ก", "วงจร", "ดิจิทัล", "ไมโคร", "สื่อสาร", "iot", "ไฟฟ้า"]):
            return RoomType.LAB_ELECTRIC.value

        return RoomType.LAB_ELECTRIC.value if course_type != CourseType.THEORY.value else RoomType.CLASSROOM.value

    @classmethod
    def process_import(
        cls,
        file_bytes: bytes,
        filename: str,
        mode: str,
        data_manager: Any
    ) -> Dict[str, Any]:
        """
        Processes bulk import of courses and lesson assignments.
        Strictly filters out any row without a valid course code.
        mode: 'replace' | 'append'
        """
        rows = cls.parse_file_to_rows(file_bytes, filename)
        if not rows:
            raise ValueError("ไม่พบข้อมูลในไฟล์ที่อัปโหลด")

        # Load existing data
        all_data = data_manager.get_all_data()
        teachers = all_data.get("teachers", [])
        rooms = all_data.get("rooms", [])
        groups = all_data.get("groups", [])
        courses = {c["id"]: c for c in all_data.get("courses", [])}
        assignments = all_data.get("assignments", [])

        if mode == "replace":
            courses = {}
            assignments = []

        imported_count = 0
        skipped_count = 0
        new_teachers_count = 0
        new_groups_count = 0
        errors = []

        for idx, row in enumerate(rows, start=2):
            code = cls._find_field(row, ["รหัสวิชา", "รหัส", "code", "course_code"])
            name = cls._find_field(row, ["ชื่อวิชา", "รายวิชา", "name", "course_name", "วิชา"])

            # Handle case where code and name might be inverted in some columns
            if not cls.is_valid_course_code(code) and cls.is_valid_course_code(name):
                code, name = name, code

            # STRICT REQUIREMENT: Only import rows with a valid course code
            # Skip category headers, section titles, summaries, notes, or empty rows
            if not cls.is_valid_course_code(code):
                skipped_count += 1
                continue

            if not name:
                name = f"วิชา {code}"

            theory_raw = cls._find_field(row, ["ท", "ทฤษฎี", "theory"])
            practice_raw = cls._find_field(row, ["ป", "ปฏิบัติ", "practice"])
            periods_raw = cls._find_field(row, ["จำนวนคาบ", "คาบ", "periods", "weekly_periods", "ชม.", "ชั่วโมง"])
            course_type_raw = cls._find_field(row, ["ประเภทวิชา", "type", "course_type"])
            room_type_raw = cls._find_field(row, ["ประเภทห้อง", "ห้อง", "room_type", "ห้องที่ต้องการ"])
            teacher_name = cls._find_field(row, ["ครูผู้สอน", "ผู้สอน", "teacher", "อาจารย์", "ครู"])
            primary_group_name = cls._find_field(row, ["กลุ่มผู้เรียน", "กลุ่มเรียน", "group", "ห้องเรียนผู้เรียน", "ห้องเรียน", "ชั้นปี"])
            secondary_group_name = cls._find_field(row, ["กลุ่มเรียนรวม", "เรียนรวม", "secondary_group", "กลุ่มร่วม"])
            rotation_raw = cls._find_field(row, ["หมุนเวียนฐาน", "หมุนฐาน", "rotation", "is_rotation"])

            # Calculate Periods
            periods = None
            if periods_raw:
                try:
                    periods = int(float(periods_raw))
                except (ValueError, TypeError):
                    pass

            if periods is None:
                t_val = 0
                p_val = 0
                has_tp = False
                try:
                    if theory_raw:
                        t_val = int(float(theory_raw))
                        has_tp = True
                except (ValueError, TypeError):
                    pass
                try:
                    if practice_raw:
                        p_val = int(float(practice_raw))
                        has_tp = True
                except (ValueError, TypeError):
                    pass

                if has_tp and (t_val + p_val) > 0:
                    periods = t_val + p_val

            if not periods or periods <= 0:
                periods = 2

            # Determine Rotation & Course Type
            is_rotation = any(w in rotation_raw.lower() for w in ["ใช่", "yes", "true", "1"])
            course_type = None

            if any(w in course_type_raw.lower() for w in ["หมุน", "rotation"]):
                course_type = CourseType.ROTATION_BASE.value
            elif any(w in course_type_raw.lower() for w in ["ปฏิบัติ", "practice", "lab"]):
                course_type = CourseType.ROTATION_BASE.value if is_rotation else CourseType.PRACTICE.value
            elif any(w in course_type_raw.lower() for w in ["ทฤษฎี", "theory", "บรรยาย"]):
                course_type = CourseType.THEORY.value
            elif is_rotation:
                course_type = CourseType.ROTATION_BASE.value
            else:
                try:
                    p_val = int(float(practice_raw)) if practice_raw else 0
                except (ValueError, TypeError):
                    p_val = 0
                if p_val > 0:
                    course_type = CourseType.PRACTICE.value
                elif periods >= 3:
                    course_type = CourseType.PRACTICE.value
                else:
                    course_type = CourseType.THEORY.value

            # Determine Room Type
            room_type = cls._determine_room_type(name, course_type, room_type_raw)

            # Match or Create Teacher
            teacher_id = ""
            if teacher_name and teacher_name != "(ยังไม่ระบุครูผู้สอน)":
                matched_t = next((t for t in teachers if t["name"].strip() == teacher_name), None)
                if matched_t:
                    teacher_id = matched_t["id"]
                else:
                    teacher_id = f"T_AUTO_{len(teachers)+1}"
                    new_t = {
                        "id": teacher_id,
                        "name": teacher_name,
                        "max_periods_per_day": 6,
                        "max_periods_per_week": 24,
                        "unavailable_slots": []
                    }
                    teachers.append(new_t)
                    new_teachers_count += 1
            else:
                unassigned = next((t for t in teachers if t["id"] == "T_UNASSIGNED"), None)
                if unassigned:
                    teacher_id = unassigned["id"]
                elif teachers:
                    teacher_id = teachers[0]["id"]
                else:
                    teacher_id = "T_DEFAULT"

            # Match or Create Primary Group
            primary_group_id = ""
            if primary_group_name:
                matched_g = next((g for g in groups if g["name"].strip() == primary_group_name or g["id"] == primary_group_name), None)
                if matched_g:
                    primary_group_id = matched_g["id"]
                else:
                    primary_group_id = f"G_AUTO_{len(groups)+1}"
                    level = "HIGH_VOC_CERT" if any(w in primary_group_name for w in ["ปวส", "ชอ.4", "ชอ.5"]) else "VOC_CERT"
                    new_g = {
                        "id": primary_group_id,
                        "name": primary_group_name,
                        "level": level,
                        "student_count": 20,
                        "pvs_18_weeks": True,
                        "is_internship": "ฝึกงาน" in primary_group_name
                    }
                    groups.append(new_g)
                    new_groups_count += 1
            else:
                sheet_name = row.get("_sheet_name", "")
                matched_g = next((g for g in groups if g["name"].strip() == sheet_name or sheet_name in g["name"]), None) if sheet_name else None
                if matched_g:
                    primary_group_id = matched_g["id"]
                elif groups:
                    primary_group_id = groups[0]["id"]
                else:
                    primary_group_id = "G_DEFAULT"

            # Match or Create Secondary Group (if any)
            secondary_group_id = None
            if secondary_group_name:
                matched_sec = next((g for g in groups if g["name"].strip() == secondary_group_name or g["id"] == secondary_group_name), None)
                if matched_sec:
                    secondary_group_id = matched_sec["id"]
                else:
                    secondary_group_id = f"G_AUTO_{len(groups)+1}"
                    level = "HIGH_VOC_CERT" if any(w in secondary_group_name for w in ["ปวส", "ชอ.4", "ชอ.5"]) else "VOC_CERT"
                    new_g = {
                        "id": secondary_group_id,
                        "name": secondary_group_name,
                        "level": level,
                        "student_count": 20,
                        "pvs_18_weeks": True,
                        "is_internship": "ฝึกงาน" in secondary_group_name
                    }
                    groups.append(new_g)
                    new_groups_count += 1

            clean_code = re.sub(r'[^0-9A-Za-z]', '_', code)
            course_id = f"C_{clean_code}_{primary_group_id}"

            # Handle 6-period sessions split (standard vocational split 3+3)
            if periods == 6:
                periods_per_session = 3
                sessions_per_week = 2
            else:
                periods_per_session = periods
                sessions_per_week = 1

            courses[course_id] = {
                "id": course_id,
                "name": name,
                "code": code,
                "course_type": course_type,
                "periods_per_session": periods_per_session,
                "sessions_per_week": sessions_per_week,
                "required_room_type": room_type,
                "allow_merge": bool(secondary_group_id),
                "base_id": f"BASE_{clean_code}" if is_rotation else None
            }

            ass_id = f"ASS_{clean_code}_{primary_group_id}"
            assignments.append({
                "id": ass_id,
                "course_id": course_id,
                "primary_group_id": primary_group_id,
                "teacher_id": teacher_id,
                "secondary_group_id": secondary_group_id,
                "is_rotation": is_rotation
            })

            imported_count += 1

        # Commit updates to data manager
        data_manager.teachers = teachers
        data_manager.groups = groups
        data_manager.courses = courses
        data_manager.assignments = assignments
        data_manager._save()

        return {
            "is_success": True,
            "imported_count": imported_count,
            "skipped_count": skipped_count,
            "new_teachers_count": new_teachers_count,
            "new_groups_count": new_groups_count,
            "total_courses": len(courses),
            "total_assignments": len(assignments),
            "errors": errors
        }

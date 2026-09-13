import tests
import unittest
import io
import openpyxl
from src.api.data_manager import data_manager
from src.api.main import _run_solver
from src.api.official_template_exporter import OfficialTemplateExporter

class TestOfficialTemplateExporter(unittest.TestCase):
    def setUp(self):
        data_manager.reset_to_default()
        self.teachers, self.rooms, self.groups, self.courses, self.assignments = data_manager.get_solver_models()
        self.res_dto = _run_solver(self.teachers, self.rooms, self.groups, self.assignments, 5, 12, 6, 5.0)
        self.schedule = [s.model_dump() for s in self.res_dto.schedule]
        self.config = data_manager.get_all_data()
        self.exporter = OfficialTemplateExporter(self.schedule, self.config)

    def test_export_single_group(self):
        group_id = self.groups[0].id
        xlsx_bytes = self.exporter.export_single_view("group", group_id, block=1)
        self.assertGreater(len(xlsx_bytes), 1000)
        
        # Verify valid openpyxl workbook
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertEqual(len(wb.sheetnames), 1)
        ws = wb.active
        self.assertEqual(ws["B9"].value, self.groups[0].name)
        self.assertEqual(ws["B6"].value, "1/2569")
        print("✅ test_export_single_group passed!")

    def test_export_single_teacher(self):
        teacher_id = self.teachers[0].id
        xlsx_bytes = self.exporter.export_single_view("teacher", teacher_id, block=1)
        self.assertGreater(len(xlsx_bytes), 1000)
        
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        self.assertEqual(ws["B8"].value, self.teachers[0].name)
        print("✅ test_export_single_teacher passed!")

    def test_export_all_groups(self):
        xlsx_bytes = self.exporter.export_all_groups(block=1)
        self.assertGreater(len(xlsx_bytes), 1000)
        
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertEqual(len(wb.sheetnames), len(self.groups))
        print("✅ test_export_all_groups passed!")

    def test_export_full_college_package(self):
        xlsx_bytes = self.exporter.export_full_college_package(block=1)
        self.assertGreater(len(xlsx_bytes), 1000)
        
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        expected_sheets = len(self.groups) + len(self.teachers) + len(self.rooms)
        self.assertEqual(len(wb.sheetnames), expected_sheets)
        print(f"✅ test_export_full_college_package passed ({len(wb.sheetnames)} sheets created)!")

    def test_api_export_official_template(self):
        from fastapi.testclient import TestClient
        from src.api.main import app
        client = TestClient(app)

        # GET single
        res = client.get("/api/export/official-template?view_type=group&view_id=G_CHO_1_1&block=1")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertGreater(len(res.content), 1000)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(len(wb.sheetnames), 1)
        print("✅ test_api_export_official_template GET passed!")

        # POST with custom schedule
        res_post = client.post("/api/export/official-template", json={
            "view_type": "all_groups",
            "block": 1,
            "schedule": self.schedule
        })
        self.assertEqual(res_post.status_code, 200)
        self.assertEqual(res_post.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb_post = openpyxl.load_workbook(io.BytesIO(res_post.content))
        self.assertEqual(len(wb_post.sheetnames), len(self.groups))
        print("✅ test_api_export_official_template POST passed!")

if __name__ == "__main__":
    unittest.main()

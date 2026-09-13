import openpyxl
import os
import json
import re
from pathlib import Path

def add_additional_curriculum_plans():
    base_dir = Path(__file__).resolve().parent.parent
    config_path = base_dir / "data" / "timetable_config.json"
    backup_path = base_dir / "data" / "timetable_config.backup_before_additional_plans.json"

    plan_dir = Path("/Users/allarmmac/Downloads/แผนการเรียน")
    files = {
        "PVC": plan_dir / "รหัส 69 ปวช.  หลักสูตร67 .xlsx",
        "PVS": plan_dir / "รหัส 69 ปวส. (ทวิ) หลักสูตร67 .xlsx",
        "PVS_M6": plan_dir / "รหัส 69 ปวส. ม.6 (ทวิ) หลักสูตร67 .xlsx"
    }

    if config_path.exists():
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        with open(backup_path, "w", encoding="utf-8") as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    else:
        config = {
            "teachers": [],
            "rooms": [],
            "groups": [],
            "courses": {},
            "assignments": []
        }

    teachers = config.get("teachers", [])
    if not any(t["id"] == "T_UNASSIGNED" for t in teachers):
        teachers.insert(0, {
            "id": "T_UNASSIGNED",
            "name": "(ยังไม่ระบุครูผู้สอน)",
            "max_periods_per_day": 12,
            "max_periods_per_week": 35,
            "is_head": False,
            "qualification": "รอมอบหมาย",
            "special_duty": "รอมอบหมาย"
        })

    rooms = config.get("rooms", [])
    existing_room_ids = {r["id"] for r in rooms}
    new_rooms = [
        {"id": "LAB_EL_1", "name": "ห้องปฏิบัติการอิเล็กทรอนิกส์ 1", "room_type": "LAB_ELECTRIC", "capacity": 40},
        {"id": "LAB_EL_2", "name": "ห้องปฏิบัติการอิเล็กทรอนิกส์ 2", "room_type": "LAB_ELECTRIC", "capacity": 40},
        {"id": "LAB_EL_3", "name": "ห้องปฏิบัติการระบบสมองกลและ IoT", "room_type": "LAB_ELECTRIC", "capacity": 40},
        {"id": "LAB_EL_4", "name": "ห้องปฏิบัติการโทรคมนาคมและเครือข่าย", "room_type": "LAB_ELECTRIC", "capacity": 40},
        {"id": "ROOM_541", "name": "ห้อง 541", "room_type": "CLASSROOM", "capacity": 35},
        {"id": "ROOM_542", "name": "ห้อง 542", "room_type": "CLASSROOM", "capacity": 35},
        {"id": "ROOM_543", "name": "ห้อง 543", "room_type": "CLASSROOM", "capacity": 35}
    ]
    for nr in new_rooms:
        if nr["id"] not in existing_room_ids:
            rooms.append(nr)
            existing_room_ids.add(nr["id"])

    groups = config.get("groups", [])
    existing_group_ids = {g["id"] for g in groups}

    new_plans = [
        (files["PVC"], "2.2570", "G_CHO_2_1", "ชอ.2/1 (ปวช.2)", "VOC_CERT", False),
        (files["PVC"], "1.2571", "G_CHO_3_1", "ชอ.3/1 (ปวช.3)", "VOC_CERT", False),
        (files["PVC"], "2.2571", "G_CHO_3_2", "ชอ.3/2 (ปวช.3 ฝึกงาน)", "VOC_CERT", True),
        (files["PVS"], "2.2570", "G_PVS_2_1", "ชอ.5/1 (ปวส.2 ทวิ ฝึกงาน)", "HIGH_VOC_CERT", True),
        (files["PVS_M6"], "2.2570", "G_PVS_M6_2_1", "ชอ.5/2 (ปวส.2 ม.6 ทวิ ฝึกงาน)", "HIGH_VOC_CERT", True)
    ]

    for fpath, sname, gid, gname, level, is_intern in new_plans:
        if gid not in existing_group_ids:
            groups.append({
                "id": gid,
                "name": gname,
                "level": level,
                "student_count": 20,
                "pvs_18_weeks": True,
                "is_internship": is_intern
            })
            existing_group_ids.add(gid)

    def determine_room_type(name, p):
        if p == 0:
            return "CLASSROOM"
        if any(k in name for k in ["ภาษา", "คณิต", "สังคม", "วิทยา", "ศีลธรรม", "พลเมือง", "สุขภาวะ", "ผู้ประกอบการ", "ธุรกิจ"]):
            return "CLASSROOM"
        if any(k in name for k in ["เชื่อม", "เครื่องมือกล", "ฝึกฝีมือ", "นิวเมติกส์"]):
            return "LAB_ENGINE"
        return "LAB_ELECTRIC"

    courses_dict = config.get("courses", {})
    assignments = config.get("assignments", [])
    existing_ass_ids = {a["id"] for a in assignments}
    ass_counter = len(assignments) + 1

    for fpath, sname, gid, gname, level, is_intern in new_plans:
        if not fpath.exists():
            print(f"Warning: file not found: {fpath}")
            continue
        wb = openpyxl.load_workbook(fpath, data_only=True)
        if sname not in wb.sheetnames:
            print(f"Warning: sheet {sname} not found in {fpath.name}")
            continue

        ws = wb[sname]
        for r in ws.iter_rows(values_only=True):
            if not r or len(r) < 5: continue
            code = str(r[0]).strip() if r[0] is not None else ""
            name = str(r[1]).strip() if r[1] is not None else ""
            t, p, c = r[2], r[3], r[4]
            if any(char.isdigit() for char in code) and "-" in code:
                try:
                    t_int = int(t) if t is not None else 0
                    p_int = int(p) if p is not None else 0
                    c_int = int(c) if c is not None else 0
                    tot = t_int + p_int
                except:
                    continue

                clean_code = re.sub(r'[^0-9A-Za-z]', '_', code)
                cid = f"C_{clean_code}"
                rtype = determine_room_type(name, p_int)

                # Rotation base logic:
                # - PVS 5-period specialized labs
                # - PVC 1.2571 3 free electives (20105-2017, 21909-2018, 20105-2027)
                is_rot = ("PVS" in gid and tot >= 5) or (code in ["20105-2017", "21909-2018", "20105-2027"])
                if is_rot:
                    ctype = "ROTATION_BASE"
                elif p_int == 0:
                    ctype = "THEORY"
                else:
                    ctype = "PRACTICE"

                if tot == 12:
                    # 12 periods project course: split into 3 sessions of 4 periods
                    for part in range(1, 4):
                        c_part_id = f"{cid}_p{part}"
                        courses_dict[c_part_id] = {
                            "id": c_part_id,
                            "name": f"{name} (ตอนที่ {part}/3)",
                            "code": code,
                            "course_type": ctype,
                            "periods_per_session": 4,
                            "sessions_per_week": 1,
                            "required_room_type": rtype,
                            "allow_merge": False,
                            "base_id": None
                        }
                        aid = f"L_{ass_counter}_{clean_code}_{gid}_p{part}"
                        if aid not in existing_ass_ids:
                            assignments.append({
                                "id": aid,
                                "course_id": c_part_id,
                                "primary_group_id": gid,
                                "teacher_id": "T_UNASSIGNED",
                                "secondary_teacher_id": None,
                                "secondary_group_id": None,
                                "is_rotation": False,
                                "teaching_mode": "SINGLE"
                            })
                            existing_ass_ids.add(aid)
                            ass_counter += 1
                elif tot == 6:
                    # 6 periods workshop/project: split into 2 sessions of 3 periods
                    for part in range(1, 3):
                        c_part_id = f"{cid}_p{part}"
                        courses_dict[c_part_id] = {
                            "id": c_part_id,
                            "name": f"{name} (ตอนที่ {part}/2)",
                            "code": code,
                            "course_type": ctype,
                            "periods_per_session": 3,
                            "sessions_per_week": 1,
                            "required_room_type": rtype,
                            "allow_merge": False,
                            "base_id": None
                        }
                        aid = f"L_{ass_counter}_{clean_code}_{gid}_p{part}"
                        if aid not in existing_ass_ids:
                            assignments.append({
                                "id": aid,
                                "course_id": c_part_id,
                                "primary_group_id": gid,
                                "teacher_id": "T_UNASSIGNED",
                                "secondary_teacher_id": None,
                                "secondary_group_id": None,
                                "is_rotation": False,
                                "teaching_mode": "SINGLE"
                            })
                            existing_ass_ids.add(aid)
                            ass_counter += 1
                else:
                    if cid not in courses_dict:
                        courses_dict[cid] = {
                            "id": cid,
                            "name": name,
                            "code": code,
                            "course_type": ctype,
                            "periods_per_session": tot,
                            "sessions_per_week": 1,
                            "required_room_type": rtype,
                            "allow_merge": False,
                            "base_id": f"BASE_{clean_code}" if is_rot else None
                        }
                    aid = f"L_{ass_counter}_{clean_code}_{gid}"
                    if aid not in existing_ass_ids:
                        assignments.append({
                            "id": aid,
                            "course_id": cid,
                            "primary_group_id": gid,
                            "teacher_id": "T_UNASSIGNED",
                            "secondary_teacher_id": None,
                            "secondary_group_id": None,
                            "is_rotation": is_rot,
                            "teaching_mode": "SINGLE"
                        })
                        existing_ass_ids.add(aid)
                        ass_counter += 1

    config["teachers"] = teachers
    config["rooms"] = rooms
    config["groups"] = groups
    config["courses"] = courses_dict
    config["assignments"] = assignments

    with open(config_path, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)

    print("Successfully added additional curriculum plans!")
    print(f"Total groups: {len(groups)}")
    print(f"Total unique courses: {len(courses_dict)}")
    print(f"Total lesson assignments: {len(assignments)}")

if __name__ == "__main__":
    add_additional_curriculum_plans()
